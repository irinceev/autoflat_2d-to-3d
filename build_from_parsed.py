# -*- coding: utf-8 -*-
"""
Второй проход пайплайна: parse_floorplan_image.parse() + легенда -> готовые файлы.

    python3 build_from_parsed.py plan.png --legend rooms_legend.csv --furniture furniture_legend.csv
"""
import os
import sys
import argparse
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle

sys.path.insert(0, os.path.dirname(__file__))
from geometry_engine import Mesh, add_door_geometry, add_window_geometry
from parse_floorplan_image import (parse, load_legend, load_furniture_legend, mask_to_rects,
                                    rects_to_m, DEFAULT_WALL_HEIGHT_M, DEFAULT_FURNITURE_HEIGHT)


def build_everything(image_path, legend_path, out_dir, furniture_path=None):
    os.makedirs(out_dir, exist_ok=True)
    p = parse(image_path)
    legend = load_legend(legend_path) if legend_path else {}
    flegend = load_furniture_legend(furniture_path) if furniture_path else {}
    H, px_per_m = p["H"], p["px_per_m"]

    for r in p["rooms"]:
        info = legend.get(r["id"], {})
        r["name"] = info.get("name", f"Комната №{r['id']}")
        r["height"] = info.get("height", DEFAULT_WALL_HEIGHT_M)
        r["flooring"] = info.get("flooring", "не указано")

    room_by_id = {r["id"]: r for r in p["rooms"]}
    for it in p["furniture"]:
        if it["id"] in flegend:
            it["height"] = flegend[it["id"]]
        owner_names = [room_by_id[rid]["name"] for rid in it["rooms"] if rid in room_by_id]
        it["room_name"] = owner_names[0] if owner_names else "не определена"

    blocked = [it for it in p["furniture"] if it["blocks_door"]]
    if blocked:
        print(f"ВНИМАНИЕ: {len(blocked)} предмет(ов) мебели перекрывают дверные проёмы: "
              f"{[it['id'] for it in blocked]} — переставь их на чертеже и пересобери.")

    def _opening_wall_height(o):
        hs = [room_by_id[rid]["height"] for rid in o["rooms"] if rid in room_by_id]
        return max(hs) if hs else DEFAULT_WALL_HEIGHT_M

    # ---------------- 3D-модель ----------------
    mesh = Mesh()
    for (x0, y0, x1, y1) in p["wall_rects_m"]:
        h = max([r["height"] for r in p["rooms"]
                 if _rect_touches_room(x0, y0, x1, y1, r, px_per_m, H)] or [DEFAULT_WALL_HEIGHT_M])
        mesh.add_box(x0, x1, y0, y1, 0.0, h, "Walls")

    # двери и окна - простая геометрия прямо в проёме
    for o in p["doors"]:
        gid = f"{int(o['cx'])}_{int(o['cy'])}"
        add_door_geometry(mesh, o, px_per_m, H, _opening_wall_height(o), group_id=gid)

    for o in p["windows"]:
        gid = f"{int(o['cx'])}_{int(o['cy'])}"
        add_window_geometry(mesh, o, px_per_m, H, _opening_wall_height(o), group_id=gid)

    for r in p["rooms"]:
        rects_px = mask_to_rects(r["mask"])
        rects_m = rects_to_m(rects_px, px_per_m, H)
        for (x0, y0, x1, y1) in rects_m:
            mesh.add_box(x0, x1, y0, y1, -0.05, 0.0, f"Floor_{r['id']}_{r['name']}")
            mesh.add_box(x0, x1, y0, y1, r["height"], r["height"] + 0.03,
                         f"Ceiling_{r['id']}_{r['name']}")

    for it in p["furniture"]:
        mesh.add_box(it["x0"], it["x1"], it["y0"], it["y1"], 0.0, it["height"],
                     f"Furniture_{it['id']}_{it['type']}_{it['room_name']}")

    obj_path = os.path.join(out_dir, "apartment_model_scanned.obj")
    mesh.write_obj(obj_path)
    print(f"3D-модель: {obj_path}")

    # ---------------- метрики ----------------
    metrics = {}
    for r in p["rooms"]:
        rid = r["id"]
        perim = _mask_perimeter_m(r["mask"], px_per_m)
        doors_here = [o for o in p["doors"] if rid in o["rooms"]]
        windows_here = [o for o in p["windows"] if rid in o["rooms"]]
        opening_area = sum(o["width_m"] * o["height_m"] for o in doors_here + windows_here)
        gross_wall = perim * r["height"]
        metrics[rid] = dict(
            name=r["name"], floor_area=r["area_m2"], perimeter=round(perim, 2),
            height=r["height"], ceiling_area=r["area_m2"],
            wall_area_gross=round(gross_wall, 2),
            wall_area_net=round(max(gross_wall - opening_area, 0), 2),
            doors=len(doors_here), windows=len(windows_here),
            door_width_sum=round(sum(o["width_m"] for o in doors_here), 2),
            flooring=r["flooring"],
        )
        metrics[rid]["skirting_length"] = round(max(perim - metrics[rid]["door_width_sum"], 0), 2)

    # ---------------- Excel ----------------
    _export_excel(metrics, p["furniture"], out_dir)

    # ---------------- визуализация 2D ----------------
    _visualize_2d(p, metrics, out_dir)

    return p, metrics


def _rect_touches_room(x0, y0, x1, y1, room, px_per_m, img_h):
    bc0, br0, bc1, br1 = room["bbox"]
    br0_m, br1_m = (img_h - br1) / px_per_m, (img_h - br0) / px_per_m
    bc0_m, bc1_m = bc0 / px_per_m, bc1 / px_per_m
    pad = 0.3
    return not (x1 < bc0_m - pad or x0 > bc1_m + pad or y1 < br0_m - pad or y0 > br1_m + pad)


def _mask_perimeter_m(mask, px_per_m):
    import cv2
    cnts, _ = cv2.findContours(mask.astype(np.uint8), cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if not cnts:
        return 0.0
    length_px = sum(cv2.arcLength(c, True) for c in cnts)
    return length_px / px_per_m


def _export_excel(metrics, furniture, out_dir):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Объёмы работ"
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(["Помещение", "Показатель", "Ед. изм.", "Количество"])
    for c in range(1, 5):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for rid in sorted(metrics):
        m = metrics[rid]
        rows = [
            (m["name"], "Площадь пола", "м²", m["floor_area"]),
            (m["name"], "Площадь стен (за вычетом проёмов)", "м²", m["wall_area_net"]),
            (m["name"], "Площадь потолка", "м²", m["ceiling_area"]),
            (m["name"], "Периметр помещения", "м", m["perimeter"]),
            (m["name"], "Количество дверей", "шт.", m["doors"]),
            (m["name"], "Количество окон", "шт.", m["windows"]),
            (m["name"], "Длина плинтуса", "м", m["skirting_length"]),
            (m["name"], "Тип напольного покрытия", "-", m["flooring"]),
        ]
        for row in rows:
            ws.append(row)
            rr = ws.max_row
            for c in range(1, 5):
                ws.cell(row=rr, column=c).border = border

    total_floor = round(sum(m["floor_area"] for m in metrics.values()), 2)
    total_wall = round(sum(m["wall_area_net"] for m in metrics.values()), 2)
    ws.append([])
    ws.append(["ИТОГО", "Площадь пола", "м²", total_floor])
    ws.append(["ИТОГО", "Площадь стен (нетто)", "м²", total_wall])
    for i, wdt in enumerate([22, 34, 10, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"

    ws3 = wb.create_sheet("Мебель")
    ws3.append(["ID", "Тип", "Комната", "Ширина, м", "Глубина, м", "Высота, м", "Перекрывает дверь?"])
    for c in range(1, 8):
        cell = ws3.cell(row=1, column=c)
        cell.fill, cell.font = header_fill, header_font
    for it in sorted(furniture, key=lambda x: x["id"]):
        ws3.append([it["id"], it["type"], it["room_name"], it["w_m"], it["d_m"],
                    it["height"], "ДА" if it["blocks_door"] else ""])
    for i, wdt in enumerate([6, 14, 20, 12, 12, 12, 18], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = wdt

    path = os.path.join(out_dir, "smeta_scanned.xlsx")
    wb.save(path)
    print(f"Excel: {path}")


def _visualize_2d(parsed, metrics, out_dir):
    fig, ax = plt.subplots(figsize=(14, 10))
    H, px_per_m = parsed["H"], parsed["px_per_m"]
    for r in parsed["rooms"]:
        for (c0, r0, c1, r1) in mask_to_rects(r["mask"]):
            x0, x1 = c0 / px_per_m, c1 / px_per_m
            y0, y1 = (H - r1) / px_per_m, (H - r0) / px_per_m
            ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, facecolor="#f0e6cc", edgecolor="none"))
        cx, cy = r["cx"] / px_per_m, (H - r["cy"]) / px_per_m
        ax.text(cx, cy, f"{r['name']}\n{r['area_m2']} м²", ha="center", va="center", fontsize=8)
    for (x0, y0, x1, y1) in parsed["wall_rects_m"]:
        ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, facecolor="#3a3a3a", edgecolor="none"))
    for it in parsed["furniture"]:
        color = "#c0392b" if it["blocks_door"] else "#8B4513"
        ax.add_patch(Rectangle((it["x0"], it["y0"]), it["x1"] - it["x0"], it["y1"] - it["y0"],
                                facecolor=color, edgecolor="#3a2210", alpha=0.85, linewidth=0.5))
        ax.text((it["x0"] + it["x1"]) / 2, (it["y0"] + it["y1"]) / 2, it["type"] or f"F{it['id']}",
                ha="center", va="center", fontsize=6, color="white")
    ax.set_aspect("equal")
    ax.autoscale()
    ax.set_title("Распознанный план (проверка результата парсинга)")
    path = os.path.join(out_dir, "floorplan_2d_scanned.png")
    plt.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    print(f"Визуализация: {path}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("image")
    ap.add_argument("--legend", required=True)
    ap.add_argument("--furniture", default=None)
    
    PROJECT_DIR = os.path.dirname(os.path.abspath(__file__)) 
    DEFAULT_OUTPUT_DIR = os.path.join(PROJECT_DIR, "outputs") 
    ap.add_argument("--out", default=DEFAULT_OUTPUT_DIR) 
    args = ap.parse_args() 
    os.makedirs(args.out, exist_ok=True)
    
    build_everything(args.image, args.legend, args.out, furniture_path=args.furniture)